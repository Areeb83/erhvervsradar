"""Read the user's spreadsheet, find the links, and write results back into it.

The results go into new columns of the *same* workbook, so the user's own
columns and formatting stay as they were. Uploading a results file again
updates those columns instead of adding them twice.
"""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from zoneinfo import ZoneInfo

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet

from .check import CheckResult

MAX_ROWS = 200
DANISH_TIME = ZoneInfo("Europe/Copenhagen")
LOOKS_LIKE_URL = re.compile(r"^(https?://|www\.)\S+$", re.I)
URL_HEADER = re.compile(r"url|link|annonce|webadresse", re.I)

TYPE_LABELS = {
    "office": "Office",
    "warehouse_industrial": "Warehouse & industrial",
    "retail": "Retail",
    "hospitality": "Hotel & restaurant",
    "mixed": "Mixed use",
    "other": "Other",
}
SOURCE_LABELS = {"edc": "EDC Erhverv", "savills": "Savills"}
STATUS_LABELS = {
    "available": "Available",
    "rented": "Rented",
    "not_found": "Listing gone",
    "unsupported_site": "Site not supported",
    "not_a_listing": "Not a listing page",
    "invalid_url": "Not a link",
    "blocked_by_robots": "Blocked by robots.txt",
    "error": "Couldn't load",
}


def result_label(result: CheckResult) -> str:
    if result.status == "ok" and result.listing:
        return STATUS_LABELS[result.listing.status]
    return STATUS_LABELS.get(result.status, STATUS_LABELS["error"])


def _listing(attr: str):
    return lambda r: getattr(r.listing, attr) if r.listing else None


# (header, value, column width, number format)
OUTPUT_COLUMNS = [
    ("Result", result_label, 18, None),
    ("Address", _listing("address"), 28, None),
    ("Postcode", _listing("zip_code"), 10, None),
    ("City", _listing("city"), 18, None),
    (
        "Property type",
        lambda r: TYPE_LABELS.get(r.listing.property_type) if r.listing else None,
        22,
        None,
    ),
    ("Size m² (from)", _listing("size_m2_min"), 14, "#,##0"),
    ("Size m² (to)", _listing("size_m2_max"), 14, "#,##0"),
    ("Rent per year kr. (from)", _listing("yearly_rent_min"), 22, "#,##0"),
    ("Rent per year kr. (to)", _listing("yearly_rent_max"), 22, "#,##0"),
    ("Rent per m² kr.", lambda r: r.listing.rent_per_m2 if r.listing else None, 16, "#,##0"),
    ("Source", lambda r: SOURCE_LABELS.get(r.listing.source) if r.listing else None, 14, None),
    (
        "Checked at (Danish time)",
        lambda r: r.checked_at.astimezone(DANISH_TIME).replace(tzinfo=None),
        22,
        "yyyy-mm-dd hh:mm",
    ),
    ("Note", lambda r: r.message, 44, None),
]


class SpreadsheetError(ValueError):
    """A problem with the uploaded file, worded for the person who uploaded it."""


@dataclass
class LinkRow:
    row: int  # row number in the sheet (1-based, like Excel)
    url: str


@dataclass
class SheetJob:
    workbook: Workbook
    sheet: Worksheet
    url_column: int
    url_header: str
    rows: list[LinkRow]
    skipped: int = 0  # rows beyond MAX_ROWS
    results: dict[int, CheckResult] = field(default_factory=dict)


# -- reading --------------------------------------------------------------------


def cell_text(cell) -> str:
    """What a link cell really points to: a clickable link's target, else its text."""
    if cell.hyperlink is not None and cell.hyperlink.target:
        return str(cell.hyperlink.target).strip()
    return "" if cell.value is None else str(cell.value).strip()


def _decode(data: bytes) -> str:
    # Danish Excel often saves CSV as Windows-1252 (æ, ø, å), not UTF-8.
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")


def read_csv(data: bytes) -> Workbook:
    text = _decode(data)
    first_line = text.splitlines()[0] if text else ""
    # Danish Excel separates columns with ";" because "," is the decimal mark.
    delimiter = ";" if first_line.count(";") > first_line.count(",") else ","
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Results"
    for row in csv.reader(io.StringIO(text), delimiter=delimiter):
        if any(field.strip() for field in row):
            sheet.append(row)
    return workbook


def read_workbook(filename: str, data: bytes) -> Workbook:
    name = filename.lower()
    if name.endswith(".csv"):
        return read_csv(data)
    if name.endswith(".xlsx"):
        try:
            return load_workbook(io.BytesIO(data))
        except Exception as exc:
            raise SpreadsheetError("This .xlsx file could not be opened. Is it damaged?") from exc
    if name.endswith(".xls"):
        raise SpreadsheetError("Old .xls files aren't supported. In Excel, save it as .xlsx first.")
    raise SpreadsheetError("Please upload an .xlsx or .csv file.")


def find_url_column(sheet: Worksheet) -> tuple[int, bool] | None:
    """(column number, has header row), or None if no column holds links."""
    counts: dict[int, int] = {}
    for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 60)):
        for cell in row:
            if LOOKS_LIKE_URL.match(cell_text(cell)):
                counts[cell.column] = counts.get(cell.column, 0) + 1
    if not counts:
        return None
    named = [c for c in counts if URL_HEADER.search(cell_text(sheet.cell(1, c)))]
    column = named[0] if named else max(counts, key=counts.get)
    has_header = not LOOKS_LIKE_URL.match(cell_text(sheet.cell(1, column)))
    return column, has_header


def load_job(filename: str, data: bytes) -> SheetJob:
    workbook = read_workbook(filename, data)
    sheet = next((ws for ws in workbook.worksheets if ws.max_row > 1 or ws["A1"].value), None)
    if sheet is None:
        raise SpreadsheetError("This file is empty.")
    found = find_url_column(sheet)
    if found is None:
        raise SpreadsheetError(
            "No links were found in this file. Put one link per row in a single column, "
            "for example under a header called “URL”."
        )
    column, has_header = found
    if not has_header:
        sheet.insert_rows(1)  # add a header row so the result columns get names
        sheet.cell(1, column, "URL")

    rows = [
        LinkRow(r, text)
        for r in range(2, sheet.max_row + 1)
        if (text := cell_text(sheet.cell(r, column)))
    ]
    return SheetJob(
        workbook=workbook,
        sheet=sheet,
        url_column=column,
        url_header=cell_text(sheet.cell(1, column)),
        rows=rows[:MAX_ROWS],
        skipped=max(0, len(rows) - MAX_ROWS),
    )


# -- writing --------------------------------------------------------------------


def write_results(job: SheetJob) -> None:
    sheet = job.sheet
    existing = {cell_text(c): c.column for c in sheet[1] if cell_text(c)}
    next_column = max(sheet.max_column, job.url_column) + 1

    for header, value, width, number_format in OUTPUT_COLUMNS:
        column = existing.get(header)
        if column is None:
            column = next_column
            next_column += 1
            cell = sheet.cell(1, column, header)
            cell.font = Font(bold=True)
            sheet.column_dimensions[cell.column_letter].width = width
        for link in job.rows:
            result = job.results.get(link.row)
            if result is None:
                continue  # stopped before this row
            cell = sheet.cell(link.row, column, value(result))
            if number_format and cell.value is not None:
                cell.number_format = number_format

    url_letter = sheet.cell(1, job.url_column).column_letter
    if (sheet.column_dimensions[url_letter].width or 0) < 30:
        sheet.column_dimensions[url_letter].width = 45


def to_bytes(workbook: Workbook) -> bytes:
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


EXAMPLE_ROWS = [
    ("Vester Voldgade 104, København V", "https://www.edc.dk/erhverv/erhvervslejemaal/kontor/1552-koebenhavn-v/vester-voldgade-104/114058334/"),
    ("Lyskær 5, Herlev", "https://www.edc.dk/erhverv/erhvervslejemaal/kontor/2730-herlev/lyskaer-5/27801563/"),
    ("Alleen 4, Skanderborg", "https://www.edc.dk/erhverv/erhvervslejemaal/detail/8660-skanderborg/alleen-4/87301724/"),
    ("Rygårds Allé 133, Hellerup", "https://search.savills.com/dk/da/ejendomsoplysninger/dff8a9c0-92ab-4c6e-afce-f1d7e327c393"),
    ("Jenagade 22, København S", "https://search.savills.com/dk/da/ejendomsoplysninger/1087c4ba-9c56-4aae-80c8-b90dcf9edab3"),
    ("A site that isn't supported", "https://www.example.com/listing/123"),
]  # fmt: skip


def example_workbook() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Listings"
    sheet.append(["Property", "URL"])
    sheet["A1"].font = sheet["B1"].font = Font(bold=True)
    for row in EXAMPLE_ROWS:
        sheet.append(row)
    sheet.column_dimensions["A"].width = 34
    sheet.column_dimensions["B"].width = 95
    return to_bytes(workbook)
