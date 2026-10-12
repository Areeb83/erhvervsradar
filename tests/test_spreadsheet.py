import io
from datetime import UTC, datetime

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill

from erhvervsradar.check import CheckResult
from erhvervsradar.models import RENTED, Listing
from erhvervsradar.spreadsheet import (
    MAX_ROWS,
    OUTPUT_COLUMNS,
    SpreadsheetError,
    example_workbook,
    load_job,
    to_bytes,
    write_results,
)

EDC = "https://www.edc.dk/erhverv/erhvervslejemaal/kontor/1552-koebenhavn-v/vester-voldgade-104/114058334/"


def xlsx(rows: list[list], hyperlinks: dict[str, str] | None = None) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Mine"
    for row in rows:
        ws.append(row)
    for ref, target in (hyperlinks or {}).items():
        ws[ref].hyperlink = target
    ws["A1"].fill = PatternFill("solid", fgColor="FFFF00")  # user formatting to keep
    return to_bytes(wb)


def ok(address="Vester Voldgade 104 2 TV", status="available") -> CheckResult:
    listing = Listing(
        source="edc", source_id="1", url=EDC, address=address, zip_code="1552",
        city="København V", property_type="office", size_m2_min=178, size_m2_max=178,
        yearly_rent_min=382_700, yearly_rent_max=382_700, status=status,
    )  # fmt: skip
    return CheckResult(
        "ok", EDC, listing=listing, checked_at=datetime(2026, 10, 12, 19, 30, tzinfo=UTC)
    )


def test_finds_the_link_column_by_header():
    job = load_job("x.xlsx", xlsx([["Name", "Link", "Notes"], ["A", EDC, "keep me"]]))
    assert job.url_column == 2 and job.url_header == "Link"
    assert [(r.row, r.url) for r in job.rows] == [(2, EDC)]


def test_finds_links_without_a_header_and_adds_one():
    job = load_job("x.xlsx", xlsx([[EDC], ["www.edc.dk/erhverv/"]]))
    assert job.sheet.cell(1, 1).value == "URL"
    assert [r.row for r in job.rows] == [2, 3]


def test_uses_the_target_of_clickable_links():
    job = load_job("x.xlsx", xlsx([["URL"], ["Vester Voldgade"]], hyperlinks={"A2": EDC}))
    assert job.rows[0].url == EDC


def test_danish_csv_with_semicolons_and_windows_encoding():
    data = f"Ejendom;URL\nKøbmagergade;{EDC}\n".encode("cp1252")
    job = load_job("liste.csv", data)
    assert job.url_header == "URL" and job.rows[0].url == EDC
    assert job.sheet.cell(2, 1).value == "Købmagergade"


def test_rejects_bad_files_with_clear_messages():
    with pytest.raises(SpreadsheetError, match="No links"):
        load_job("x.csv", b"name,city\nA,B\n")
    with pytest.raises(SpreadsheetError, match="xls"):
        load_job("old.xls", b"\xd0\xcf")
    with pytest.raises(SpreadsheetError, match="could not be opened"):
        load_job("broken.xlsx", b"not a zip")


def test_caps_the_number_of_rows():
    job = load_job("x.csv", ("URL\n" + f"{EDC}\n" * (MAX_ROWS + 5)).encode())
    assert len(job.rows) == MAX_ROWS and job.skipped == 5


def test_writes_results_next_to_the_users_columns():
    job = load_job(
        "x.xlsx", xlsx([["Name", "URL", "Notes"], ["A", EDC, "keep me"], ["B", "nope", ""]])
    )
    job.results = {
        2: ok(),
        3: CheckResult("invalid_url", "nope", message="This is not a valid web address."),
    }
    write_results(job)
    ws = load_workbook(io.BytesIO(to_bytes(job.workbook))).active

    headers = [c.value for c in ws[1]]
    assert headers[:3] == ["Name", "URL", "Notes"]  # untouched
    assert headers[3:] == [name for name, *_ in OUTPUT_COLUMNS]
    assert ws["C2"].value == "keep me"
    assert ws["A1"].fill.fgColor.rgb.endswith("FFFF00")  # user formatting kept

    row = {h: c.value for h, c in zip(headers, ws[2], strict=True)}
    assert row["Result"] == "Available"
    assert row["Rent per year kr. (from)"] == 382_700
    assert ws.cell(2, headers.index("Rent per year kr. (from)") + 1).number_format == "#,##0"
    assert row["Checked at (Danish time)"] == datetime(2026, 10, 12, 21, 30)  # UTC+2 in October
    assert ws.cell(3, headers.index("Result") + 1).value == "Not a link"


def test_reuploading_a_results_file_updates_instead_of_duplicating():
    job = load_job("x.xlsx", xlsx([["URL"], [EDC]]))
    job.results = {2: ok()}
    write_results(job)
    again = load_job("x (checked).xlsx", to_bytes(job.workbook))
    again.results = {2: ok(status=RENTED)}
    write_results(again)
    ws = again.sheet
    assert ws.max_column == 1 + len(OUTPUT_COLUMNS)
    assert ws.cell(2, 2).value == "Rented"


def test_example_sheet_loads():
    job = load_job("example.xlsx", example_workbook())
    assert len(job.rows) == 6 and job.url_header == "URL"
